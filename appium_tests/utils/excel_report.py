"""
DentAI Appium — Excel & HTML Report Generator
===============================================
Reads test results collected by conftest.py and generates:
  - appium/reports/excel/Appium_E2E_Test_Report.xlsx  (5 sheets)
  - appium/reports/html/Appium_E2E_Test_Report.html

Usage (called automatically at end of pytest session via conftest.py):
    from appium_tests.utils.excel_report import generate_reports
    generate_reports(results_list)
"""

import datetime
import os
import platform
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    raise ImportError("openpyxl not installed. Run: pip install openpyxl==3.1.5")

from appium_tests.config.config import EXCEL_DIR, HTML_DIR, SCREENSHOTS_DIR
from appium_tests.utils.test_data import (
    DEVICE_NAME, PLATFORM_VERSION, APP_VERSION, APP_PACKAGE
)

# ── Colour palette (8-char ARGB for openpyxl) ─────────────────────────────────
C = {
    "hdr_dark":  "FF0D1B2A",
    "hdr_blue":  "FF1A3A5C",
    "pass_bg":   "FFD4EDDA",
    "pass_fg":   "FF155724",
    "fail_bg":   "FFF8D7DA",
    "fail_fg":   "FF721C24",
    "skip_bg":   "FFFFF3CD",
    "skip_fg":   "FF856404",
    "info_bg":   "FFD6EAF8",
    "info_fg":   "FF1A3A5C",
    "row_even":  "FFF0F4F8",
    "row_odd":   "FFFFFFFF",
    "white":     "FFFFFFFF",
    "border":    "FFCCCCCC",
}


def _fill(c):
    return PatternFill("solid", fgColor=c)


def _bdr():
    s = Side(style="thin", color=C["border"])
    return Border(left=s, right=s, top=s, bottom=s)


def _font(bold=False, color="FF2C3E50", size=10):
    return Font(name="Calibri", size=size, bold=bold, color=color)


def _align(h="left", wrap=False):
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap)


def _set_col(ws, ci, width):
    ws.column_dimensions[get_column_letter(ci)].width = width


def _hdr_row(ws, headers, widths, row=1, bg=None):
    bg = bg or C["hdr_dark"]
    for ci, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row, ci, h)
        c.font      = _font(True, C["white"], 10)
        c.fill      = _fill(bg)
        c.border    = _bdr()
        c.alignment = _align("center")
        _set_col(ws, ci, w)
    ws.row_dimensions[row].height = 22
    ws.freeze_panes = f"A{row+1}"


def _status_colors(status: str):
    if status == "PASS":
        return C["pass_bg"], C["pass_fg"]
    if status == "SKIP":
        return C["skip_bg"], C["skip_fg"]
    return C["fail_bg"], C["fail_fg"]


# ─────────────────────────────────────────────────────────────────────────────
# PUBLIC ENTRY POINT
# ─────────────────────────────────────────────────────────────────────────────

def generate_reports(results: list[dict]) -> tuple[Path, Path]:
    """
    Generate Excel and HTML reports from the collected test results.

    Each result dict must contain:
        tc_id, name, module, status, start_time, end_time,
        duration, error, screenshot, precondition, steps, expected, actual

    Returns:
        Tuple of (excel_path, html_path).
    """
    now      = datetime.datetime.utcnow()
    now_str  = now.strftime("%Y-%m-%d %H:%M:%S UTC")
    excel_path = EXCEL_DIR / "Appium_E2E_Test_Report.xlsx"
    html_path  = HTML_DIR  / "Appium_E2E_Test_Report.html"

    _generate_excel(results, now_str, excel_path)
    _generate_html(results, now_str, html_path)

    return excel_path, html_path


# ─────────────────────────────────────────────────────────────────────────────
# EXCEL GENERATION
# ─────────────────────────────────────────────────────────────────────────────

def _generate_excel(results: list[dict], now_str: str, out: Path):
    wb = openpyxl.Workbook()

    _sheet1_test_summary(wb, results, now_str)
    _sheet2_execution_summary(wb, results, now_str)
    _sheet3_module_analysis(wb, results)
    _sheet4_detailed_results(wb, results)
    _sheet5_device_info(wb, now_str)

    wb.save(str(out))
    print(f"SUCCESS Excel report -> {out}  ({out.stat().st_size:,} bytes)")


def _sheet1_test_summary(wb, results, now_str):
    ws = wb.active
    ws.title = "Test Summary"

    # Title banner
    ws.merge_cells("A1:L1")
    tc = ws.cell(1, 1, "DentAI Mobile — Appium E2E Test Summary")
    tc.font = _font(True, C["white"], 13)
    tc.fill = _fill(C["hdr_dark"])
    tc.alignment = _align("center")
    ws.row_dimensions[1].height = 32

    headers = [
        "TC ID", "Test Name", "Module", "Status",
        "Start Time", "End Time", "Duration (s)",
        "Error Message", "Screenshot", "Device", "Android Version"
    ]
    widths = [12, 35, 18, 8, 20, 20, 13, 45, 40, 20, 16]
    _hdr_row(ws, headers, widths, row=2, bg=C["hdr_blue"])
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}2"

    for ri, r in enumerate(results, 3):
        bg, fg = _status_colors(r.get("status", "FAIL"))
        row_bg = C["row_even"] if ri % 2 == 0 else C["row_odd"]
        vals = [
            r.get("tc_id", ""),
            r.get("name", ""),
            r.get("module", ""),
            r.get("status", "FAIL"),
            r.get("start_time", ""),
            r.get("end_time", ""),
            f"{r.get('duration', 0):.2f}",
            r.get("error", "")[:120],
            r.get("screenshot", ""),
            DEVICE_NAME,
            f"Android {PLATFORM_VERSION}",
        ]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(ri, ci, val)
            c.border    = _bdr()
            c.alignment = _align("left" if ci in (2, 8, 9) else "center", wrap=ci == 8)
            if ci == 4:  # Status
                c.font = _font(True, fg, 10)
                c.fill = _fill(bg)
            else:
                c.font = _font(size=9)
                c.fill = _fill(row_bg)
        ws.row_dimensions[ri].height = 18


def _sheet2_execution_summary(wb, results, now_str):
    ws = wb.create_sheet("Execution Summary")

    ws.merge_cells("A1:B1")
    tc = ws.cell(1, 1, "DentAI Mobile — Appium Execution Summary")
    tc.font = _font(True, C["white"], 13)
    tc.fill = _fill(C["hdr_dark"])
    tc.alignment = _align("center")
    ws.row_dimensions[1].height = 30
    ws.cell(1, 2).fill = _fill(C["hdr_dark"])

    total   = len(results)
    passed  = sum(1 for r in results if r.get("status") == "PASS")
    failed  = sum(1 for r in results if r.get("status") == "FAIL")
    skipped = sum(1 for r in results if r.get("status") == "SKIP")
    pass_pct = (passed / max(total, 1)) * 100
    fail_pct = (failed / max(total, 1)) * 100
    total_dur = sum(r.get("duration", 0) for r in results)
    mins, secs = divmod(int(total_dur), 60)

    _hdr_row(ws, ["Metric", "Value"], [32, 40], row=2, bg=C["hdr_blue"])

    rows = [
        ("Total Test Cases",    str(total)),
        ("Passed",              str(passed)),
        ("Failed",              str(failed)),
        ("Skipped",             str(skipped)),
        ("Pass Percentage",     f"{pass_pct:.2f}%"),
        ("Fail Percentage",     f"{fail_pct:.2f}%"),
        ("Total Execution Time", f"{mins}m {secs}s"),
        ("Device",              DEVICE_NAME),
        ("Android Version",     f"Android {PLATFORM_VERSION}"),
        ("App Version",         APP_VERSION),
        ("App Package",         APP_PACKAGE),
        ("Execution Date",      now_str),
    ]

    for ri, (k, v) in enumerate(rows, 3):
        row_bg = C["row_even"] if ri % 2 == 0 else C["row_odd"]
        kc = ws.cell(ri, 1, k)
        vc = ws.cell(ri, 2, v)
        kc.font = _font(True, "FF2C3E50", 10)
        kc.fill = _fill(row_bg)
        kc.border = _bdr()
        kc.alignment = _align()

        if k == "Pass Percentage":
            bg, fg = C["pass_bg"], C["pass_fg"]
        elif k == "Fail Percentage":
            bg, fg = (C["fail_bg"], C["fail_fg"]) if failed > 0 else (C["pass_bg"], C["pass_fg"])
        else:
            bg, fg = row_bg, "FF2C3E50"
        vc.font = _font(True, fg, 10)
        vc.fill = _fill(bg)
        vc.border = _bdr()
        vc.alignment = _align()
        ws.row_dimensions[ri].height = 20

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 40

    # Overall verdict
    vri = len(rows) + 4
    ws.merge_cells(f"A{vri}:B{vri}")
    verdict = "SUCCESS  OVERALL: PASS" if failed == 0 else f"FAILED  OVERALL: FAIL — {failed} test(s) failed"
    bg = C["pass_bg"] if failed == 0 else C["fail_bg"]
    fg = C["pass_fg"] if failed == 0 else C["fail_fg"]
    vc = ws.cell(vri, 1, verdict)
    vc.font = _font(True, fg, 13)
    vc.fill = _fill(bg)
    vc.alignment = _align("center")
    vc.border = _bdr()
    ws.row_dimensions[vri].height = 30
    ws.cell(vri, 2).fill = _fill(bg)
    ws.cell(vri, 2).border = _bdr()


def _sheet3_module_analysis(wb, results):
    ws = wb.create_sheet("Module Analysis")

    ws.merge_cells("A1:G1")
    tc = ws.cell(1, 1, "DentAI Mobile — Module-wise Test Analysis")
    tc.font = _font(True, C["white"], 12)
    tc.fill = _fill(C["hdr_dark"])
    tc.alignment = _align("center")
    ws.row_dimensions[1].height = 28

    headers = ["Module", "Total", "Passed", "Failed", "Skipped", "Pass %", "Fail %"]
    widths  = [22, 8, 8, 8, 8, 10, 10]
    _hdr_row(ws, headers, widths, row=2, bg=C["hdr_blue"])

    # Aggregate by module
    modules: dict = {}
    for r in results:
        m = r.get("module", "Unknown")
        if m not in modules:
            modules[m] = {"total": 0, "passed": 0, "failed": 0, "skipped": 0}
        modules[m]["total"]   += 1
        s = r.get("status", "FAIL")
        if s == "PASS":   modules[m]["passed"]  += 1
        elif s == "SKIP": modules[m]["skipped"] += 1
        else:             modules[m]["failed"]  += 1

    for ri, (mod, d) in enumerate(sorted(modules.items()), 3):
        row_bg = C["row_even"] if ri % 2 == 0 else C["row_odd"]
        pass_pct = d["passed"] / max(d["total"], 1) * 100
        fail_pct = d["failed"] / max(d["total"], 1) * 100
        vals = [mod, d["total"], d["passed"], d["failed"], d["skipped"],
                f"{pass_pct:.1f}%", f"{fail_pct:.1f}%"]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(ri, ci, val)
            c.border = _bdr()
            c.alignment = _align("left" if ci == 1 else "center")
            if ci == 3:   # passed
                c.font = _font(True, C["pass_fg"], 10)
                c.fill = _fill(C["pass_bg"])
            elif ci == 4 and d["failed"] > 0:  # failed
                c.font = _font(True, C["fail_fg"], 10)
                c.fill = _fill(C["fail_bg"])
            else:
                c.font = _font(size=10)
                c.fill = _fill(row_bg)
        ws.row_dimensions[ri].height = 20


def _sheet4_detailed_results(wb, results):
    ws = wb.create_sheet("Detailed Results")

    ws.merge_cells("A1:J1")
    tc = ws.cell(1, 1, "DentAI Mobile — Detailed Test Results")
    tc.font = _font(True, C["white"], 12)
    tc.fill = _fill(C["hdr_dark"])
    tc.alignment = _align("center")
    ws.row_dimensions[1].height = 28

    headers = [
        "TC ID", "Test Name", "Precondition", "Test Steps",
        "Expected Result", "Actual Result", "Status",
        "Duration (s)", "Error", "Screenshot Path"
    ]
    widths = [12, 30, 25, 50, 35, 35, 8, 12, 50, 45]
    _hdr_row(ws, headers, widths, row=2, bg=C["hdr_blue"])
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}2"

    for ri, r in enumerate(results, 3):
        bg, fg = _status_colors(r.get("status", "FAIL"))
        row_bg = C["row_even"] if ri % 2 == 0 else C["row_odd"]
        vals = [
            r.get("tc_id", ""),
            r.get("name", ""),
            r.get("precondition", ""),
            r.get("steps", ""),
            r.get("expected", ""),
            r.get("actual", r.get("error", ""))[:200],
            r.get("status", "FAIL"),
            f"{r.get('duration', 0):.2f}",
            r.get("error", "")[:150],
            r.get("screenshot", ""),
        ]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(ri, ci, val)
            c.border    = _bdr()
            c.alignment = _align("left", wrap=ci in (4, 5, 6, 9))
            if ci == 7:
                c.font = _font(True, fg, 10)
                c.fill = _fill(bg)
            else:
                c.font = _font(size=9)
                c.fill = _fill(row_bg)
        ws.row_dimensions[ri].height = 40 if any(len(str(r.get(k, ""))) > 60 for k in ("steps", "expected", "error")) else 20


def _sheet5_device_info(wb, now_str):
    ws = wb.create_sheet("Device Information")

    ws.merge_cells("A1:B1")
    tc = ws.cell(1, 1, "DentAI Mobile — Test Environment Information")
    tc.font = _font(True, C["white"], 12)
    tc.fill = _fill(C["hdr_dark"])
    tc.alignment = _align("center")
    ws.row_dimensions[1].height = 28

    _hdr_row(ws, ["Property", "Value"], [30, 45], row=2, bg=C["hdr_blue"])

    import appium as _appium_pkg
    import pytest as _pytest

    try:
        appium_ver = _appium_pkg.__version__
    except Exception:
        appium_ver = "unknown"

    rows = [
        ("Device Name",      DEVICE_NAME),
        ("Android Version",  f"Android {PLATFORM_VERSION}"),
        ("Platform",         "Android"),
        ("App Package",      APP_PACKAGE),
        ("App Version",      APP_VERSION),
        ("Appium Version",   appium_ver),
        ("Python Version",   sys.version.split()[0]),
        ("Pytest Version",   _pytest.__version__),
        ("Execution Date",   now_str),
        ("Report Generated", datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("OS",               platform.system() + " " + platform.release()),
    ]

    for ri, (k, v) in enumerate(rows, 3):
        row_bg = C["row_even"] if ri % 2 == 0 else C["row_odd"]
        for ci, val in enumerate((k, v), 1):
            c = ws.cell(ri, ci, val)
            c.font = _font(True if ci == 1 else False, "FF2C3E50", 10)
            c.fill = _fill(row_bg)
            c.border = _bdr()
            c.alignment = _align()
        ws.row_dimensions[ri].height = 20

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 45


# ─────────────────────────────────────────────────────────────────────────────
# HTML GENERATION
# ─────────────────────────────────────────────────────────────────────────────

def _generate_html(results: list[dict], now_str: str, out: Path):
    total   = len(results)
    passed  = sum(1 for r in results if r.get("status") == "PASS")
    failed  = sum(1 for r in results if r.get("status") == "FAIL")
    skipped = sum(1 for r in results if r.get("status") == "SKIP")
    pass_pct = (passed / max(total, 1)) * 100
    total_dur = sum(r.get("duration", 0) for r in results)
    mins, secs = divmod(int(total_dur), 60)
    result_label = "PASS" if failed == 0 else "FAIL"
    res_color = "#155724" if failed == 0 else "#721c24"
    res_bg    = "#d4edda" if failed == 0 else "#f8d7da"

    def badge(status):
        bg = {"PASS": "#d4edda", "FAIL": "#f8d7da", "SKIP": "#fff3cd"}.get(status, "#f8d7da")
        fg = {"PASS": "#155724", "FAIL": "#721c24", "SKIP": "#856404"}.get(status, "#721c24")
        return f'<span style="background:{bg};color:{fg};padding:2px 8px;border-radius:3px;font-weight:bold;font-size:11px">{status}</span>'

    # Module table rows
    modules: dict = {}
    for r in results:
        m = r.get("module", "Unknown")
        if m not in modules:
            modules[m] = {"total": 0, "passed": 0, "failed": 0, "skipped": 0}
        modules[m]["total"] += 1
        s = r.get("status", "FAIL")
        if s == "PASS":   modules[m]["passed"] += 1
        elif s == "SKIP": modules[m]["skipped"] += 1
        else:             modules[m]["failed"] += 1

    mod_rows = ""
    for mod, d in sorted(modules.items()):
        pp = d["passed"] / max(d["total"], 1) * 100
        mod_rows += f"""<tr>
            <td>{mod}</td>
            <td style="text-align:center">{d['total']}</td>
            <td style="text-align:center;color:#155724;font-weight:bold">{d['passed']}</td>
            <td style="text-align:center;color:{'#721c24' if d['failed'] else '#155724'};font-weight:bold">{d['failed']}</td>
            <td style="text-align:center">{d['skipped']}</td>
            <td style="text-align:center">{pp:.1f}%</td>
        </tr>\n"""

    # Individual test rows
    test_rows = ""
    for r in results:
        ss = r.get("screenshot", "")
        ss_link = f'<a href="file://{ss}" target="_blank">📷</a>' if ss else "—"
        err = r.get("error", "")[:120]
        test_rows += f"""<tr>
            <td>{r.get('tc_id','')}</td>
            <td>{r.get('name','')}</td>
            <td>{r.get('module','')}</td>
            <td style="text-align:center">{badge(r.get('status','FAIL'))}</td>
            <td style="text-align:right">{r.get('duration',0):.2f}s</td>
            <td style="font-size:11px;color:#dc2626">{err}</td>
            <td style="text-align:center">{ss_link}</td>
        </tr>\n"""

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>DentAI — Appium E2E Test Report</title>
<style>
  *{{box-sizing:border-box;margin:0;padding:0}}
  body{{font-family:'Segoe UI',Arial,sans-serif;background:#f4f6f9;color:#1e293b}}
  .header{{background:linear-gradient(135deg,#0d1b2a,#1a3a5c);color:#fff;padding:32px 40px}}
  .header h1{{font-size:24px;margin-bottom:6px}}
  .header p{{font-size:13px;opacity:.75}}
  .container{{max-width:1400px;margin:24px auto;padding:0 24px}}
  .result-badge{{display:inline-block;padding:8px 24px;border-radius:6px;font-size:15px;font-weight:700;margin:16px 0;background:{res_bg};color:{res_color};border:2px solid {res_color}}}
  .cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(140px,1fr));gap:14px;margin:20px 0}}
  .card{{background:#fff;border-radius:8px;padding:18px 16px;text-align:center;box-shadow:0 2px 8px rgba(0,0,0,.08)}}
  .card .val{{font-size:28px;font-weight:700;margin:8px 0}}
  .card .lbl{{font-size:11px;color:#64748b;text-transform:uppercase;letter-spacing:.5px}}
  .card.green .val{{color:#15803d}}.card.red .val{{color:#dc2626}}.card.blue .val{{color:#1d4ed8}}.card.amber .val{{color:#b45309}}
  .section{{background:#fff;border-radius:8px;padding:22px;margin:20px 0;box-shadow:0 2px 8px rgba(0,0,0,.08)}}
  .section h2{{font-size:16px;color:#0d1b2a;margin-bottom:16px;border-bottom:2px solid #3b82f6;padding-bottom:8px}}
  table{{width:100%;border-collapse:collapse}}
  th{{background:#0d1b2a;color:#fff;padding:10px 8px;text-align:left;font-size:12px}}
  td{{padding:9px 8px;border-bottom:1px solid #e2e8f0;font-size:12px}}
  tr:nth-child(even){{background:#f8fafc}}
  footer{{text-align:center;padding:24px;font-size:11px;color:#94a3b8}}
  .mobile-icon{{font-size:32px;margin-bottom:8px}}
</style>
</head>
<body>
<div class="header">
  <p class="mobile-icon">📱</p>
  <h1>DentAI Mobile — Appium E2E Test Report</h1>
  <p>Android Automated Testing · DentAI CBCT AI Platform</p>
  <p>Generated: {now_str} · Device: {DEVICE_NAME} · Android {PLATFORM_VERSION}</p>
</div>

<div class="container">
  <div class="result-badge">{'SUCCESS' if failed==0 else 'FAILED'} OVERALL RESULT: {result_label}</div>

  <div class="cards">
    <div class="card blue"><div class="val">{total}</div><div class="lbl">Total Tests</div></div>
    <div class="card green"><div class="val">{passed}</div><div class="lbl">Passed</div></div>
    <div class="card {'red' if failed else 'green'}"><div class="val">{failed}</div><div class="lbl">Failed</div></div>
    <div class="card amber"><div class="val">{skipped}</div><div class="lbl">Skipped</div></div>
    <div class="card {'green' if pass_pct>=90 else 'amber'}"><div class="val">{pass_pct:.1f}%</div><div class="lbl">Pass Rate</div></div>
    <div class="card blue"><div class="val">{mins}m{secs}s</div><div class="lbl">Duration</div></div>
  </div>

  <div class="section">
    <h2>Test Configuration</h2>
    <table>
      <tr><th>Property</th><th>Value</th></tr>
      <tr><td>Device</td><td>{DEVICE_NAME}</td></tr>
      <tr><td>Android Version</td><td>Android {PLATFORM_VERSION}</td></tr>
      <tr><td>App Package</td><td>{APP_PACKAGE}</td></tr>
      <tr><td>App Version</td><td>{APP_VERSION}</td></tr>
      <tr><td>Total Tests</td><td>{total}</td></tr>
      <tr><td>Execution Time</td><td>{mins}m {secs}s</td></tr>
      <tr><td>Date</td><td>{now_str}</td></tr>
    </table>
  </div>

  <div class="section">
    <h2>Module Analysis</h2>
    <table>
      <tr><th>Module</th><th>Total</th><th>Passed</th><th>Failed</th><th>Skipped</th><th>Pass %</th></tr>
      {mod_rows}
    </table>
  </div>

  <div class="section">
    <h2>Individual Test Results</h2>
    <table>
      <tr><th>TC ID</th><th>Test Name</th><th>Module</th><th>Status</th><th>Duration</th><th>Error</th><th>Screenshot</th></tr>
      {test_rows}
    </table>
  </div>

</div>
<footer>DentAI Appium E2E Report · Generated {now_str}</footer>
</body>
</html>"""

    out.write_text(html, encoding="utf-8")
    print(f"SUCCESS HTML report  -> {out}  ({out.stat().st_size:,} bytes)")
