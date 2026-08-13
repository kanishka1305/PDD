"""
excel_reporter.py
=================
Core Excel report engine for the DentAI automation framework.

Produces:
  reports/excel/<safe_test_name>.xlsx   — one file per test
  reports/excel/overall_test_report.xlsx — aggregated summary

Called exclusively by the pytest plugin in automation/plugin/conftest_excel.py.
Never call this directly from test files.

All colors use 8-character ARGB format required by openpyxl (FF prefix).
"""

from __future__ import annotations

import os
import platform
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter


# ── Output directory ──────────────────────────────────────────────────────────
REPORTS_DIR = Path(__file__).parent.parent.parent / "reports" / "excel"
SCREENSHOTS_DIR = Path(__file__).parent.parent.parent / "reports" / "screenshots"


def _ensure_dirs() -> None:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    SCREENSHOTS_DIR.mkdir(parents=True, exist_ok=True)


# ── ARGB color palette ────────────────────────────────────────────────────────
C = {
    "header_dark":   "FF1A252F",   # dark navy — header background
    "header_pass":   "FF155724",   # dark green
    "header_fail":   "FF721C24",   # dark red
    "header_skip":   "FF856404",   # dark amber
    "header_env":    "FF0D47A1",   # dark blue
    "header_stats":  "FF4A235A",   # dark purple
    "pass_bg":       "FFD4EDDA",   # light green
    "fail_bg":       "FFF8D7DA",   # light red
    "skip_bg":       "FFFFF3CD",   # light amber
    "err_bg":        "FFFCE4EC",   # light pink
    "row_even":      "FFECF0F1",
    "row_odd":       "FFFFFFFF",
    "white":         "FFFFFFFF",
    "border":        "FFCCCCCC",
    "pass_text":     "FF155724",
    "fail_text":     "FF721C24",
    "skip_text":     "FF856404",
}


# ── Style helpers ─────────────────────────────────────────────────────────────

def _fill(argb: str) -> PatternFill:
    return PatternFill("solid", fgColor=argb)


def _font(bold=False, color="FF2C3E50", size=10, italic=False) -> Font:
    return Font(name="Calibri", size=size, bold=bold,
                color=color, italic=italic)


def _border() -> Border:
    s = Side(style="thin", color=C["border"])
    return Border(left=s, right=s, top=s, bottom=s)


def _align(h="left", wrap=True) -> Alignment:
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap)


def _status_colors(status: str) -> tuple[str, str]:
    """Return (bg_argb, text_argb) for a status string."""
    s = status.upper()
    if s == "PASS":
        return C["pass_bg"], C["pass_text"]
    if s == "FAIL":
        return C["fail_bg"], C["fail_text"]
    if s in ("SKIP", "SKIPPED"):
        return C["skip_bg"], C["skip_text"]
    return C["err_bg"], C["fail_text"]   # ERROR / UNKNOWN


def _write_header_row(ws, headers: list[str], bg: str = C["header_dark"],
                      row: int = 1) -> None:
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row, col, text)
        cell.font = _font(bold=True, color=C["white"], size=10)
        cell.fill = _fill(bg)
        cell.border = _border()
        cell.alignment = _align("center", wrap=False)
    ws.row_dimensions[row].height = 22


def _auto_col_widths(ws, widths: list[int]) -> None:
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def _safe_name(nodeid: str) -> str:
    """Convert a pytest nodeid to a filesystem-safe filename (no extension)."""
    # automation/tests/test_authentication.py::TestAuthentication::test_auth_001_...
    # → test_auth_001_...
    parts = nodeid.replace("\\", "/").split("::")
    name = parts[-1] if parts else nodeid
    # Strip parametrize suffix like [param0]
    name = re.sub(r"\[.*\]$", "", name)
    # Replace unsafe chars
    name = re.sub(r"[^a-zA-Z0-9_\-]", "_", name)
    # Collapse multiple underscores
    name = re.sub(r"_+", "_", name).strip("_")
    return name[:80]  # limit length


def _get_env_info() -> dict[str, str]:
    """Collect environment metadata once."""
    try:
        import selenium
        sel_ver = selenium.__version__
    except Exception:
        sel_ver = "unknown"

    try:
        import pytest
        pytest_ver = pytest.__version__
    except Exception:
        pytest_ver = "unknown"

    # Chrome version — read from env var set by the CI step or detect locally
    chrome_ver = os.environ.get("CHROME_VERSION", "")
    if not chrome_ver:
        try:
            import subprocess
            out = subprocess.check_output(
                ["google-chrome", "--version"], stderr=subprocess.DEVNULL,
                timeout=5
            ).decode().strip()
            chrome_ver = out.split()[-1] if out else "unknown"
        except Exception:
            try:
                import subprocess
                out = subprocess.check_output(
                    ["google-chrome-stable", "--version"], stderr=subprocess.DEVNULL,
                    timeout=5
                ).decode().strip()
                chrome_ver = out.split()[-1] if out else "unknown"
            except Exception:
                chrome_ver = "unknown"

    return {
        "Python Version":   sys.version.split()[0],
        "Selenium Version": sel_ver,
        "Pytest Version":   pytest_ver,
        "Chrome Version":   chrome_ver,
        "Operating System": platform.platform(),
        "Environment":      "GitHub Actions" if os.environ.get("GITHUB_ACTIONS") else "Local",
        "Execution Date":   datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
        "Repository":       os.environ.get("GITHUB_REPOSITORY", "kanishka1305/PDD"),
        "Branch":           os.environ.get("GITHUB_REF_NAME", "local"),
        "Run Number":       os.environ.get("GITHUB_RUN_NUMBER", "local"),
    }


# ═════════════════════════════════════════════════════════════════════════════
# INDIVIDUAL TEST REPORT
# ═════════════════════════════════════════════════════════════════════════════

def write_individual_report(record: dict) -> Path:
    """
    Write one .xlsx file for a single test execution.

    record keys (all strings unless noted):
        nodeid, name, file, function, module, status,
        start_time (datetime), end_time (datetime), duration (float),
        browser, os, environment, error_message, screenshot_path
    """
    _ensure_dirs()

    safe = _safe_name(record.get("nodeid", record.get("name", "unknown")))
    out_path = REPORTS_DIR / f"{safe}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Result"

    status = (record.get("status") or "UNKNOWN").upper()
    bg, fg = _status_colors(status)

    # ── Section 1: Test Information ───────────────────────────────────────────
    info_headers = ["Field", "Value"]
    _write_header_row(ws, info_headers, bg=C["header_dark"], row=1)

    start_dt: Any = record.get("start_time")
    end_dt: Any   = record.get("end_time")
    start_str = start_dt.strftime("%Y-%m-%d %H:%M:%S") if isinstance(start_dt, datetime) else str(start_dt or "")
    end_str   = end_dt.strftime("%Y-%m-%d %H:%M:%S")   if isinstance(end_dt,   datetime) else str(end_dt   or "")
    duration  = record.get("duration", 0.0)

    fields = [
        ("Test ID",         record.get("nodeid",      "")),
        ("Test Name",       record.get("name",        "")),
        ("Test File",       record.get("file",        "")),
        ("Test Function",   record.get("function",    "")),
        ("Module",          record.get("module",      "")),
        ("Status",          status),
        ("Start Time",      start_str),
        ("End Time",        end_str),
        ("Duration (s)",    f"{float(duration):.3f}"),
        ("Browser",         record.get("browser",     "Chrome")),
        ("Operating System",record.get("os",          platform.system())),
        ("Environment",     record.get("environment", "GitHub Actions" if os.environ.get("GITHUB_ACTIONS") else "Local")),
        ("Error Message",   record.get("error_message", "")),
        ("Screenshot Path", record.get("screenshot_path", "")),
    ]

    for row_idx, (field, value) in enumerate(fields, 2):
        # Field name cell
        fc = ws.cell(row_idx, 1, field)
        fc.font      = _font(bold=True, size=10)
        fc.fill      = _fill(C["row_even"] if row_idx % 2 == 0 else C["row_odd"])
        fc.border    = _border()
        fc.alignment = _align("left", wrap=False)

        # Value cell
        vc = ws.cell(row_idx, 2, value)
        vc.border    = _border()
        vc.alignment = _align("left", wrap=True)

        # Special formatting for Status row
        if field == "Status":
            vc.font = _font(bold=True, color=fg, size=11)
            vc.fill = _fill(bg)
            fc.fill = _fill(bg)
        elif field == "Error Message" and value:
            vc.font = _font(color=C["fail_text"], size=9, italic=True)
            vc.fill = _fill(C["fail_bg"])
        else:
            vc.font = _font(size=10)
            vc.fill = _fill(C["row_even"] if row_idx % 2 == 0 else C["row_odd"])

        ws.row_dimensions[row_idx].height = 20 if field != "Error Message" else 40

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 80
    ws.freeze_panes = "A2"

    # ── Section 2: Environment (below the test info, with a gap) ─────────────
    gap_row = len(fields) + 3
    ws.cell(gap_row, 1, "Environment").font = _font(bold=True, size=11,
                                                    color=C["white"])
    ws.cell(gap_row, 1).fill   = _fill(C["header_env"])
    ws.cell(gap_row, 2).fill   = _fill(C["header_env"])
    ws.cell(gap_row, 1).border = _border()
    ws.cell(gap_row, 2).border = _border()
    ws.row_dimensions[gap_row].height = 20

    env = _get_env_info()
    for i, (k, v) in enumerate(env.items(), gap_row + 1):
        kc = ws.cell(i, 1, k)
        vc = ws.cell(i, 2, v)
        row_bg = C["row_even"] if i % 2 == 0 else C["row_odd"]
        for c in (kc, vc):
            c.fill = _fill(row_bg); c.border = _border()
            c.alignment = _align("left", wrap=False)
        kc.font = _font(bold=True, size=9)
        vc.font = _font(size=9)
        ws.row_dimensions[i].height = 18

    wb.save(str(out_path))
    return out_path


# ═════════════════════════════════════════════════════════════════════════════
# OVERALL REPORT
# ═════════════════════════════════════════════════════════════════════════════

def write_overall_report(records: list[dict], session_start: datetime,
                         session_end: datetime) -> Path:
    """
    Write overall_test_report.xlsx with 4 sheets:
      1. Test Summary
      2. Statistics
      3. Failed Tests
      4. Environment
    """
    _ensure_dirs()
    out_path = REPORTS_DIR / "overall_test_report.xlsx"

    wb = openpyxl.Workbook()
    _sheet_test_summary(wb, records)
    _sheet_statistics(wb, records, session_start, session_end)
    _sheet_failed_tests(wb, records)
    _sheet_environment(wb)

    # Remove the default blank sheet if it was not renamed
    default = wb.get_sheet_by_name("Sheet") if "Sheet" in wb.sheetnames else None
    if default:
        wb.remove(default)

    wb.save(str(out_path))
    return out_path


# ── Sheet 1: Test Summary ─────────────────────────────────────────────────────

def _sheet_test_summary(wb: openpyxl.Workbook, records: list[dict]) -> None:
    ws = wb.active
    ws.title = "Test Summary"

    headers = [
        "Test ID", "Test Name", "Test File", "Test Function",
        "Status", "Start Time", "End Time", "Duration (s)",
        "Browser", "Environment", "Error Message", "Screenshot",
    ]
    widths = [45, 45, 35, 35, 10, 20, 20, 12, 12, 18, 55, 50]

    _write_header_row(ws, headers, bg=C["header_dark"], row=1)
    _auto_col_widths(ws, widths)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for row_idx, r in enumerate(records, 2):
        status = (r.get("status") or "UNKNOWN").upper()
        bg, fg  = _status_colors(status)
        row_bg  = bg  # color entire row by status

        start_dt: Any = r.get("start_time")
        end_dt: Any   = r.get("end_time")
        start_str = start_dt.strftime("%Y-%m-%d %H:%M:%S") if isinstance(start_dt, datetime) else str(start_dt or "")
        end_str   = end_dt.strftime("%Y-%m-%d %H:%M:%S")   if isinstance(end_dt,   datetime) else str(end_dt   or "")

        values = [
            r.get("nodeid",          ""),
            r.get("name",            ""),
            r.get("file",            ""),
            r.get("function",        ""),
            status,
            start_str,
            end_str,
            f"{float(r.get('duration', 0)):.3f}",
            r.get("browser",         "Chrome"),
            r.get("environment",     ""),
            r.get("error_message",   ""),
            r.get("screenshot_path", ""),
        ]

        for col, val in enumerate(values, 1):
            cell = ws.cell(row_idx, col, val)
            cell.fill      = _fill(row_bg)
            cell.border    = _border()
            cell.alignment = _align("left", wrap=False)
            cell.font      = _font(size=9)

        # Bold, colored status cell
        sc = ws.cell(row_idx, 5)
        sc.font = _font(bold=True, color=fg, size=9)
        sc.alignment = _align("center", wrap=False)

        ws.row_dimensions[row_idx].height = 18


# ── Sheet 2: Statistics ───────────────────────────────────────────────────────

def _sheet_statistics(wb: openpyxl.Workbook, records: list[dict],
                      session_start: datetime, session_end: datetime) -> None:
    ws = wb.create_sheet("Statistics")

    total   = len(records)
    passed  = sum(1 for r in records if r.get("status", "").upper() == "PASS")
    failed  = sum(1 for r in records if r.get("status", "").upper() == "FAIL")
    skipped = sum(1 for r in records if r.get("status", "").upper() in ("SKIP", "SKIPPED"))
    errors  = sum(1 for r in records if r.get("status", "").upper() == "ERROR")

    pass_pct = f"{passed / total * 100:.1f}%" if total else "0%"
    fail_pct = f"{(failed + errors) / total * 100:.1f}%" if total else "0%"
    duration = (session_end - session_start).total_seconds() if session_start and session_end else 0.0

    _write_header_row(ws, ["Metric", "Value"], bg=C["header_stats"])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.freeze_panes = "A2"

    stats = [
        ("Total Tests",           total),
        ("Passed",                passed),
        ("Failed",                failed),
        ("Skipped",               skipped),
        ("Errors",                errors),
        ("Pass Percentage",       pass_pct),
        ("Fail Percentage",       fail_pct),
        ("Total Execution Time",  f"{duration:.2f}s"),
        ("Session Start",         session_start.strftime("%Y-%m-%d %H:%M:%S") if session_start else ""),
        ("Session End",           session_end.strftime("%Y-%m-%d %H:%M:%S")   if session_end   else ""),
    ]

    for row_idx, (metric, value) in enumerate(stats, 2):
        row_bg = C["row_even"] if row_idx % 2 == 0 else C["row_odd"]

        mc = ws.cell(row_idx, 1, metric)
        mc.font = _font(bold=True, size=11); mc.fill = _fill(row_bg)
        mc.border = _border(); mc.alignment = _align("left", wrap=False)

        vc = ws.cell(row_idx, 2, value)
        vc.border = _border(); vc.alignment = _align("center", wrap=False)
        vc.font = _font(size=11)

        # Color value cells for key metrics
        if metric == "Passed":
            vc.font = _font(bold=True, color=C["pass_text"], size=12)
            vc.fill = _fill(C["pass_bg"])
        elif metric in ("Failed", "Errors"):
            vc.font = _font(bold=True, color=C["fail_text"], size=12)
            vc.fill = _fill(C["fail_bg"])
        elif metric == "Skipped":
            vc.font = _font(bold=True, color=C["skip_text"], size=12)
            vc.fill = _fill(C["skip_bg"])
        elif metric == "Pass Percentage":
            vc.font = _font(bold=True, color=C["pass_text"], size=12)
            vc.fill = _fill(C["pass_bg"])
        elif metric == "Fail Percentage":
            vc.font = _font(bold=True, color=C["fail_text"], size=12)
            vc.fill = _fill(C["fail_bg"])
        else:
            vc.fill = _fill(row_bg)

        ws.row_dimensions[row_idx].height = 24


# ── Sheet 3: Failed Tests ─────────────────────────────────────────────────────

def _sheet_failed_tests(wb: openpyxl.Workbook, records: list[dict]) -> None:
    ws = wb.create_sheet("Failed Tests")

    headers = [
        "Test Name", "Test File", "Error Message",
        "Failure Reason", "Screenshot", "Execution Time (s)",
    ]
    widths = [45, 35, 70, 70, 50, 18]

    _write_header_row(ws, headers, bg=C["header_fail"])
    _auto_col_widths(ws, widths)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    failed = [r for r in records if r.get("status", "").upper() in ("FAIL", "ERROR")]

    if not failed:
        nc = ws.cell(2, 1, "No failures — all tests passed!")
        nc.font = _font(bold=True, color=C["pass_text"], size=11)
        nc.fill = _fill(C["pass_bg"]); nc.border = _border()
        ws.row_dimensions[2].height = 24
        return

    for row_idx, r in enumerate(failed, 2):
        err  = r.get("error_message",   "") or ""
        vals = [
            r.get("name",            ""),
            r.get("file",            ""),
            err[:200],
            err[:300],
            r.get("screenshot_path", ""),
            f"{float(r.get('duration', 0)):.3f}",
        ]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row_idx, col, val)
            cell.fill      = _fill(C["fail_bg"])
            cell.border    = _border()
            cell.alignment = _align("left", wrap=True)
            cell.font      = _font(size=9, color=C["fail_text"])
        ws.row_dimensions[row_idx].height = 35


# ── Sheet 4: Environment ──────────────────────────────────────────────────────

def _sheet_environment(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Environment")

    _write_header_row(ws, ["Property", "Value"], bg=C["header_env"])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60
    ws.freeze_panes = "A2"

    env = _get_env_info()
    for row_idx, (k, v) in enumerate(env.items(), 2):
        row_bg = C["row_even"] if row_idx % 2 == 0 else C["row_odd"]

        kc = ws.cell(row_idx, 1, k)
        kc.font = _font(bold=True, size=10); kc.fill = _fill(row_bg)
        kc.border = _border(); kc.alignment = _align("left", wrap=False)

        vc = ws.cell(row_idx, 2, v)
        vc.font = _font(size=10); vc.fill = _fill(row_bg)
        vc.border = _border(); vc.alignment = _align("left", wrap=False)

        ws.row_dimensions[row_idx].height = 20
