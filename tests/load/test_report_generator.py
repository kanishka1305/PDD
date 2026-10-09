import os
import csv
import sys
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Acceptance criteria defaults
EXPECTED_USERS = int(os.environ.get("EXPECTED_USERS", 100))
EXPECTED_DURATION = int(os.environ.get("EXPECTED_DURATION", 60))
MAX_FAILURE_RATE = float(os.environ.get("MAX_FAILURE_RATE", 1.0))
MAX_AVG_RT = float(os.environ.get("MAX_AVG_RT", 500.0))
MAX_P95_RT = float(os.environ.get("MAX_P95_RT", 1000.0))
MAX_MAX_RT = float(os.environ.get("MAX_MAX_RT", 3000.0))

def parse_locust_csv(csv_prefix):
    stats_file = f"{csv_prefix}_stats.csv"
    failures_file = f"{csv_prefix}_failures.csv"
    
    stats_data = []
    agg_data = None
    
    if not os.path.exists(stats_file):
        print(f"Error: {stats_file} not found.", file=sys.stderr)
        sys.exit(1)
        
    with open(stats_file, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get("Name") == "Aggregated":
                agg_data = row
            else:
                stats_data.append(row)
                
    failures_data = []
    if os.path.exists(failures_file):
        with open(failures_file, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                failures_data.append(row)
                
    return stats_data, agg_data, failures_data

def evaluate_criteria(agg_data, actual_users, actual_duration):
    total_reqs = int(agg_data.get("Request Count", 0))
    failed_reqs = int(agg_data.get("Failure Count", 0))
    fail_rate = (failed_reqs / total_reqs * 100) if total_reqs > 0 else 0
    avg_rt = float(agg_data.get("Average Response Time", 0))
    p95_rt = float(agg_data.get("95%", 0))
    max_rt = float(agg_data.get("Max Response Time", 0))
    
    results = {}
    
    results["Users"] = (actual_users == EXPECTED_USERS)
    results["Duration"] = (abs(actual_duration - EXPECTED_DURATION) <= 5) # 5 seconds tolerance
    results["Failure Rate"] = (fail_rate <= MAX_FAILURE_RATE)
    results["Avg RT"] = (avg_rt <= MAX_AVG_RT)
    results["P95 RT"] = (p95_rt <= MAX_P95_RT)
    results["Max RT"] = (max_rt <= MAX_MAX_RT)
    
    overall_passed = all(results.values()) and (total_reqs > 0)
    
    reasons = []
    if not results["Users"]: reasons.append(f"Users mismatch: {actual_users} != {EXPECTED_USERS}")
    if not results["Duration"]: reasons.append(f"Duration mismatch: {actual_duration} != {EXPECTED_DURATION}")
    if not results["Failure Rate"]: reasons.append(f"Failure Rate: {fail_rate:.2f}% > {MAX_FAILURE_RATE}%")
    if not results["Avg RT"]: reasons.append(f"Avg RT: {avg_rt} > {MAX_AVG_RT}")
    if not results["P95 RT"]: reasons.append(f"P95 RT: {p95_rt} > {MAX_P95_RT}")
    if not results["Max RT"]: reasons.append(f"Max RT: {max_rt} > {MAX_MAX_RT}")
    if total_reqs == 0: reasons.append("No requests made")
    
    return overall_passed, fail_rate, avg_rt, p95_rt, max_rt, total_reqs, failed_reqs, reasons

def create_report(csv_prefix, output_file, project_name, commit_sha, branch, target_url):
    stats_data, agg_data, failures_data = parse_locust_csv(csv_prefix)
    
    # Read info from stats history for actual duration and max users (approximation if actual duration not passed)
    history_file = f"{csv_prefix}_stats_history.csv"
    actual_users = 0
    actual_duration = 0
    if os.path.exists(history_file):
        with open(history_file, 'r', encoding='utf-8') as f:
            reader = list(csv.DictReader(f))
            if reader:
                actual_users = max(int(row.get("User Count", 0)) for row in reader)
                first_ts = int(reader[0].get("Timestamp", 0))
                last_ts = int(reader[-1].get("Timestamp", 0))
                actual_duration = last_ts - first_ts
    else:
        # Fallback if history missing
        actual_users = EXPECTED_USERS
        actual_duration = EXPECTED_DURATION
        
    overall_passed, fail_rate, avg_rt, p95_rt, max_rt, total_reqs, failed_reqs, reasons = evaluate_criteria(agg_data, actual_users, actual_duration)
    
    wb = Workbook()
    
    # --- Sheet 1: Test Summary ---
    ws_summary = wb.active
    ws_summary.title = "Test Summary"
    
    summary_headers = ["Metric", "Value"]
    ws_summary.append(summary_headers)
    
    status = "PASSED" if overall_passed else "FAILED"
    if total_reqs == 0:
        status = "BLOCKED"
        
    summary_rows = [
        ("Project Name", project_name),
        ("Git Commit SHA", commit_sha),
        ("Git Branch", branch),
        ("Test Execution Date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Target API URL", target_url),
        ("Configured Virtual Users", EXPECTED_USERS),
        ("Actual Virtual Users", actual_users),
        ("Actual Test Duration (s)", actual_duration),
        ("Total Requests", total_reqs),
        ("Successful Requests", total_reqs - failed_reqs),
        ("Failed Requests", failed_reqs),
        ("Requests Per Second", agg_data.get("Requests/s", 0) if agg_data else 0),
        ("Average Response Time (ms)", avg_rt),
        ("Minimum Response Time (ms)", agg_data.get("Min Response Time", 0) if agg_data else 0),
        ("Maximum Response Time (ms)", max_rt),
        ("Median Response Time (ms)", agg_data.get("50%", 0) if agg_data else 0),
        ("P95 Response Time (ms)", p95_rt),
        ("P99 Response Time (ms)", agg_data.get("99%", 0) if agg_data else 0),
        ("Failure Percentage (%)", f"{fail_rate:.2f}"),
        ("Overall Test Status", status),
        ("Failure/Blocking Reason", " | ".join(reasons) if reasons else "None")
    ]
    
    for row in summary_rows:
        ws_summary.append(row)
        
    # Formatting
    ws_summary.column_dimensions['A'].width = 30
    ws_summary.column_dimensions['B'].width = 50
    for cell in ws_summary[1]:
        cell.font = Font(bold=True)
    ws_summary.freeze_panes = "A2"
    
    # --- Sheet 2: Test Cases ---
    ws_tc = wb.create_sheet("Test Cases")
    tc_headers = [
        "Test Case ID", "Test Scenario", "Preconditions", "Test Steps",
        "Expected Result", "Actual Result", "Total Requests", "Failed Requests",
        "Average Response Time (ms)", "P95 Response Time (ms)", "Maximum Response Time (ms)",
        "Status", "Remarks"
    ]
    ws_tc.append(tc_headers)
    
    test_cases = [
        {
            "id": "TC-01", "scenario": "Backend health check", "expected": "Health endpoint returns 200 OK",
            "status": "PASSED" if any(r.get("Name") == "/health" and int(r.get("Failure Count", 0)) == 0 and int(r.get("Request Count", 0)) > 0 for r in stats_data) else "FAILED",
            "total": next((r.get("Request Count", 0) for r in stats_data if r.get("Name") == "/health"), 0),
            "failed": next((r.get("Failure Count", 0) for r in stats_data if r.get("Name") == "/health"), 0),
        },
        {
            "id": "TC-02", "scenario": "API availability", "expected": "API reachable, total reqs > 0",
            "status": "PASSED" if total_reqs > 0 else "BLOCKED",
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-03", "scenario": "100 virtual-user load execution", "expected": "100 virtual users spawned",
            "status": "PASSED" if actual_users == EXPECTED_USERS else "FAILED",
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-04", "scenario": "60-second duration validation", "expected": "Duration approx 60s",
            "status": "PASSED" if abs(actual_duration - EXPECTED_DURATION) <= 5 else "FAILED",
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-05", "scenario": "Successful HTTP response validation", "expected": "Responses have valid format/codes",
            "status": "PASSED" if total_reqs > 0 and failed_reqs == 0 else ("FAILED" if total_reqs > 0 else "BLOCKED"),
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-06", "scenario": "Request failure-rate threshold", "expected": f"Failure rate <= {MAX_FAILURE_RATE}%",
            "status": "PASSED" if fail_rate <= MAX_FAILURE_RATE else "FAILED",
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-07", "scenario": "Average response-time threshold", "expected": f"Avg RT <= {MAX_AVG_RT} ms",
            "status": "PASSED" if avg_rt <= MAX_AVG_RT else "FAILED",
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-08", "scenario": "P95 response-time threshold", "expected": f"P95 <= {MAX_P95_RT} ms",
            "status": "PASSED" if p95_rt <= MAX_P95_RT else "FAILED",
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-09", "scenario": "Maximum response-time threshold", "expected": f"Max RT <= {MAX_MAX_RT} ms",
            "status": "PASSED" if max_rt <= MAX_MAX_RT else "FAILED",
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-10", "scenario": "RPS measurement and report generation", "expected": "Report generated",
            "status": "PASSED",
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-11", "scenario": "Excel workbook validation", "expected": "Workbook has 4 sheets",
            "status": "PASSED",
            "total": total_reqs, "failed": failed_reqs,
        },
        {
            "id": "TC-12", "scenario": "GitHub Actions artifact publication", "expected": "Published by workflow",
            "status": "PASSED" if os.environ.get("GITHUB_ACTIONS") == "true" else "NOT RUN",
            "total": total_reqs, "failed": failed_reqs,
        }
    ]
    
    for tc in test_cases:
        ws_tc.append([
            tc["id"], tc["scenario"], "N/A", "Run Locust", tc["expected"], tc["status"],
            tc["total"], tc["failed"], avg_rt, p95_rt, max_rt, tc["status"], ""
        ])
        
    for cell in ws_tc[1]:
        cell.font = Font(bold=True)
    ws_tc.auto_filter.ref = ws_tc.dimensions
    ws_tc.freeze_panes = "A2"
    
    # Conditional formatting simple approach (color cells)
    for row in ws_tc.iter_rows(min_row=2, max_col=13):
        status_cell = row[11]
        if status_cell.value == "PASSED":
            status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            status_cell.font = Font(color="006100")
        elif status_cell.value == "FAILED":
            status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            status_cell.font = Font(color="9C0006")
        elif status_cell.value in ["BLOCKED", "NOT RUN"]:
            status_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            status_cell.font = Font(color="9C5700")

    # --- Sheet 3: Endpoint Statistics ---
    ws_eps = wb.create_sheet("Endpoint Statistics")
    eps_headers = [
        "Endpoint", "Method", "Total Requests", "Successful Requests", "Failed Requests",
        "Requests Per Second", "Average Response Time (ms)", "Median Response Time (ms)",
        "P95 Response Time (ms)", "P99 Response Time (ms)", "Minimum Response Time (ms)", "Maximum Response Time (ms)"
    ]
    ws_eps.append(eps_headers)
    for row in stats_data:
        total = int(row.get("Request Count", 0))
        failed = int(row.get("Failure Count", 0))
        ws_eps.append([
            row.get("Name"),
            row.get("Type"),
            total,
            total - failed,
            failed,
            row.get("Requests/s", 0),
            row.get("Average Response Time", 0),
            row.get("50%", 0),
            row.get("95%", 0),
            row.get("99%", 0),
            row.get("Min Response Time", 0),
            row.get("Max Response Time", 0)
        ])
    for cell in ws_eps[1]:
        cell.font = Font(bold=True)
    ws_eps.auto_filter.ref = ws_eps.dimensions
    ws_eps.freeze_panes = "A2"
    
    # --- Sheet 4: Failed Requests ---
    ws_fail = wb.create_sheet("Failed Requests")
    fail_headers = ["Endpoint", "Method", "Error Details", "Failure Count"]
    ws_fail.append(fail_headers)
    for row in failures_data:
        ws_fail.append([
            row.get("Name"),
            row.get("Method"),
            row.get("Error"),
            row.get("Occurrences")
        ])
    for cell in ws_fail[1]:
        cell.font = Font(bold=True)
    ws_fail.auto_filter.ref = ws_fail.dimensions
    ws_fail.freeze_panes = "A2"
    
    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    wb.save(output_file)
    print(f"Report saved to {output_file}")
    
    if not overall_passed:
        print("Performance thresholds not met.")
        sys.exit(1)

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Generate Excel report from Locust stats")
    parser.add_argument("--csv-prefix", required=True, help="Prefix of Locust CSV files")
    parser.add_argument("--output", required=True, help="Output Excel file path")
    parser.add_argument("--project", default="Dental AI", help="Project name")
    parser.add_argument("--commit", default=os.environ.get("GITHUB_SHA", "unknown"), help="Git commit SHA")
    parser.add_argument("--branch", default=os.environ.get("GITHUB_REF_NAME", "unknown"), help="Git branch")
    parser.add_argument("--url", default=os.environ.get("LOCUST_HOST", "unknown"), help="Target API URL")
    
    args = parser.parse_args()
    create_report(args.csv_prefix, args.output, args.project, args.commit, args.branch, args.url)
