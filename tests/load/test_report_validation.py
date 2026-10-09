import os
import sys
sys.path.insert(0, os.path.dirname(__file__))
import csv
import unittest
from openpyxl import load_workbook
import tempfile
from test_report_generator import create_report

class TestReportGenerator(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.csv_prefix = os.path.join(self.temp_dir.name, "locust_stats")
        self.output_excel = os.path.join(self.temp_dir.name, "report.xlsx")
        
        # Create mock CSV files
        stats_data = [
            {"Type": "GET", "Name": "/health", "Request Count": "150", "Failure Count": "0", "Median Response Time": "100", "Average Response Time": "105", "Min Response Time": "80", "Max Response Time": "200", "Requests/s": "2.5", "50%": "100", "95%": "150", "99%": "180"},
            {"Type": "", "Name": "Aggregated", "Request Count": "150", "Failure Count": "0", "Median Response Time": "100", "Average Response Time": "105", "Min Response Time": "80", "Max Response Time": "200", "Requests/s": "2.5", "50%": "100", "95%": "150", "99%": "180"}
        ]
        
        with open(f"{self.csv_prefix}_stats.csv", 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=stats_data[0].keys())
            writer.writeheader()
            writer.writerows(stats_data)
            
        history_data = [
            {"Timestamp": "1000", "User Count": "100"},
            {"Timestamp": "1060", "User Count": "100"}
        ]
        with open(f"{self.csv_prefix}_stats_history.csv", 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=history_data[0].keys())
            writer.writeheader()
            writer.writerows(history_data)
            
        failures_data = []
        with open(f"{self.csv_prefix}_failures.csv", 'w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=["Method", "Name", "Error", "Occurrences"])
            writer.writeheader()
            writer.writerows(failures_data)
            
        os.environ["EXPECTED_USERS"] = "100"
        os.environ["EXPECTED_DURATION"] = "60"

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_generator_creates_workbook(self):
        # Override exit since it calls sys.exit if it fails
        try:
            create_report(self.csv_prefix, self.output_excel, "Test Proj", "abc1234", "main", "http://test")
        except SystemExit:
            pass
            
        self.assertTrue(os.path.exists(self.output_excel))
        
        wb = load_workbook(self.output_excel)
        sheet_names = wb.sheetnames
        self.assertIn("Test Summary", sheet_names)
        self.assertIn("Test Cases", sheet_names)
        self.assertIn("Endpoint Statistics", sheet_names)
        self.assertIn("Failed Requests", sheet_names)
        
        ws_summary = wb["Test Summary"]
        # Basic check of values
        found_status = False
        for row in ws_summary.iter_rows(values_only=True):
            if row[0] == "Overall Test Status":
                self.assertEqual(row[1], "PASSED")
                found_status = True
        self.assertTrue(found_status)

if __name__ == "__main__":
    unittest.main()
